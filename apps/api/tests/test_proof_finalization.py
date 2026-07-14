"""Shared finalization boundary + block-typed document evidence regressions.

Covers the canonical sequence for every proof type:

    owner → owned project → owned proof → retained artifact → analysis →
    normalized evidence → proof-project relationship → project-skill claims →
    claim-evidence links → report eligibility

and the block-aware document pipeline (exact table/chart/diagram/code/metric
citations instead of "document text mentions skill").
"""

from __future__ import annotations

import io
from uuid import uuid4

import pytest

from app.services.canonical_evidence_service import (
    CanonicalEvidenceConflictError,
    CanonicalEvidenceNotFoundError,
    document_evidence_object_has_locator,
    finalize_proof_evidence,
)
from app.services.optional_evidence_service import (
    OptionalEvidenceService,
    analyze_from_extracted_document,
    extract_document_text,
)
from tests.test_vbr_project_defense import USER_ID

OTHER_USER = "00000000-0000-0000-0000-00000000beef"


@pytest.fixture()
def mem_store() -> dict:
    return {}


def _project(store: dict, *, title: str, user_id: str = USER_ID, claimed: list[str] | None = None) -> dict:
    pid = str(uuid4())
    row = {
        "id": pid,
        "user_id": user_id,
        "title": title,
        "repo_url": "https://github.com/veribridge/veribridge",
        "repo_full_name": "veribridge/veribridge",
        "status": "draft",
        "metadata": {"claimed_skills": claimed or ["Python", "FastAPI"], "attached_proofs": {}},
    }
    store.setdefault("vbr_projects", {})[pid] = row
    return row


def _document(
    store: dict,
    *,
    user_id: str = USER_ID,
    with_block_locators: bool = True,
) -> dict:
    doc_id = str(uuid4())
    objects = [
        {
            "skill_name": "FastAPI",
            "snippet": "The API layer is implemented with FastAPI endpoints for scoring.",
            "reason": "Supported by a table in the 'Architecture' section referencing fastapi.",
            "section_label": "Architecture" if with_block_locators else "Normal",
            "block_type": "table" if with_block_locators else None,
            "block_index": 7 if with_block_locators else None,
            "table_cells": [["Component", "Technology"], ["API", "FastAPI"]] if with_block_locators else [],
            "page_number": None,
        },
        {
            "skill_name": "Data Analysis",
            "snippet": "Data analysis matters.",
            "reason": "Supported by document text mentioning data analysis.",
            "section_label": "Normal",
            "block_type": None,
            "page_number": None,
        },
    ]
    row = {
        "id": doc_id,
        "user_id": user_id,
        "proof_session_id": None,
        "source_type": "document",
        "status": "analyzed",
        "file_path": "VeriBridge_Rich_Document_Proof_Test.docx",
        "analysis_json": {"title": "VeriBridge Rich Document Proof Test"},
        "evidence_objects": objects,
    }
    store.setdefault("optional_evidence_submissions", {})[doc_id] = row
    artifact_id = str(uuid4())
    store.setdefault("proof_artifacts", {})[artifact_id] = {
        "id": artifact_id,
        "owner_user_id": user_id,
        "proof_type": "document",
        "artifact_type": "document_original",
        "proof_id": doc_id,
        "project_id": None,
        "retained": True,
        "file_name": "VeriBridge_Rich_Document_Proof_Test.docx",
    }
    return row


# ── Document finalization ──────────────────────────────────────────────────────


def test_document_finalization_attaches_links_and_stamps_artifact(mem_store: dict) -> None:
    project = _project(mem_store, title="VeriBridge")
    document = _document(mem_store)

    result = finalize_proof_evidence(
        mem_store,
        user_id=USER_ID,
        proof_type="document",
        proof_id=document["id"],
        project_id=project["id"],
    )

    assert result["project_relationship"]["state"] == "directly_linked"
    assert result["report_eligibility"] == "report_ready"
    # attached_proofs.documents carries the canonical edge the reports read.
    docs = project["metadata"]["attached_proofs"]["documents"]
    assert [d["document_evidence_id"] for d in docs] == [document["id"]]
    # Normalized relationship row exists.
    rels = list(mem_store["proof_project_relationships"].values())
    assert len(rels) == 1 and rels[0]["relationship_state"] == "directly_linked"
    assert rels[0]["proof_type"] == "document"
    # The retained original artifact is stamped with the project.
    artifact = next(iter(mem_store["proof_artifacts"].values()))
    assert artifact["project_id"] == project["id"]
    # Locator-backed skill counted; keyword-only skill honestly excluded.
    links = list(mem_store["vbr_claim_evidence_links"].values())
    by_status = {link["link_status"] for link in links}
    assert by_status == {"counted", "excluded"}
    counted = [l for l in links if l["link_status"] == "counted"]
    assert len(counted) == 1
    assert result["supported_skill_count"] == 1
    assert any("keyword" in w.lower() for w in result["warnings"])


def test_document_finalization_is_idempotent(mem_store: dict) -> None:
    project = _project(mem_store, title="VeriBridge")
    document = _document(mem_store)
    first = finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="document",
        proof_id=document["id"], project_id=project["id"],
    )
    second = finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="document",
        proof_id=document["id"], project_id=project["id"],
    )
    assert first["claim_link_count"] == second["claim_link_count"]
    assert len(mem_store["proof_project_relationships"]) == 1
    assert len(project["metadata"]["attached_proofs"]["documents"]) == 1
    assert len(mem_store["vbr_claim_evidence_links"]) == first["claim_link_count"]


def test_document_finalization_rejects_foreign_project(mem_store: dict) -> None:
    foreign = _project(mem_store, title="Foreign", user_id=OTHER_USER)
    document = _document(mem_store)
    with pytest.raises(CanonicalEvidenceNotFoundError):
        finalize_proof_evidence(
            mem_store, user_id=USER_ID, proof_type="document",
            proof_id=document["id"], project_id=foreign["id"],
        )
    assert not mem_store.get("proof_project_relationships")


def test_document_finalization_rejects_foreign_document(mem_store: dict) -> None:
    project = _project(mem_store, title="VeriBridge")
    document = _document(mem_store, user_id=OTHER_USER)
    with pytest.raises(CanonicalEvidenceNotFoundError):
        finalize_proof_evidence(
            mem_store, user_id=USER_ID, proof_type="document",
            proof_id=document["id"], project_id=project["id"],
        )


def test_document_finalization_rejects_second_project(mem_store: dict) -> None:
    veribridge = _project(mem_store, title="VeriBridge")
    boston = _project(mem_store, title="Boston Smart Accident Risk Rerouting")
    document = _document(mem_store)
    finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="document",
        proof_id=document["id"], project_id=veribridge["id"],
    )
    with pytest.raises(CanonicalEvidenceConflictError):
        finalize_proof_evidence(
            mem_store, user_id=USER_ID, proof_type="document",
            proof_id=document["id"], project_id=boston["id"],
        )
    # Boston never gains the document.
    assert "documents" not in (boston["metadata"].get("attached_proofs") or {})


# ── GitHub finalization ────────────────────────────────────────────────────────


def _github_proof(store: dict, *, user_id: str = USER_ID, line_level: bool = True) -> dict:
    gid = str(uuid4())
    snapshot = {}
    if line_level:
        snapshot = {
            "skill_code_evidence": [
                # A REAL implementation body: counts (strong grade + concrete purpose).
                {
                    "skill": "Python",
                    "file_path": "api.py",
                    "line_start": 252,
                    "line_end": 255,
                    "code_snippet": (
                        "@app.post(\"/predict\")\n"
                        "def predict(payload: PredictIn):\n"
                        "    preds = model.predict(payload.features)\n"
                        "    return JSONResponse({\"risk\": preds.tolist()})"
                    ),
                },
                # A sliced bare-name import fragment: line-level but UNVALIDATABLE
                # (insufficient context) — must stay context-only, never counted.
                {
                    "skill": "Data Analysis",
                    "file_path": "train.py",
                    "line_start": 22,
                    "line_end": 22,
                    "code_snippet": "    accuracy_score, f1_score, classification_report",
                },
            ]
        }
    row = {
        "id": gid,
        "user_id": user_id,
        "proof_session_id": None,
        "repo_url": "https://github.com/veribridge/veribridge",
        "repo_owner": "veribridge",
        "repo_name": "veribridge",
        "default_branch": "main",
        "status": "analyzed",
        "detected_skills": ["Python", "Docker"],
        "submitted_skill_claims": ["Python"],
        "public_safe_summary": "GitHub proof for veribridge/veribridge.",
        "repo_metadata": {},
        "analysis_snapshot": snapshot,
        "created_at": "2026-07-12T00:00:00+00:00",
        "updated_at": "2026-07-12T00:00:00+00:00",
    }
    store.setdefault("github_proof_submissions", {})[gid] = row
    return row


def test_github_finalization_counts_line_level_and_groups_repo_context(mem_store: dict) -> None:
    project = _project(mem_store, title="VeriBridge", claimed=["Python", "Docker"])
    proof = _github_proof(mem_store)
    result = finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="github",
        proof_id=proof["id"], project_id=project["id"],
    )
    assert result["project_relationship"]["state"] == "directly_linked"
    links = list(mem_store["vbr_claim_evidence_links"].values())
    counted = [l for l in links if l["link_status"] == "counted"]
    context = [l for l in links if l["link_status"] == "excluded"]
    # Python has a VALIDATED implementation body at the cited lines; Docker is
    # repo-level detection only; the Data Analysis row is a sliced import
    # fragment whose purpose cannot be determined — line-level but NEVER counted.
    assert len(counted) == 1 and counted[0]["evidence_quality"] == "primary"
    counted_skills = {l["skill_name"] for l in counted} if counted and "skill_name" in counted[0] else None
    if counted_skills is not None:
        assert counted_skills == {"Python"}
    assert len(context) == 2
    assert any("repo" in w.lower() for w in result["warnings"])
    assert any("could not be validated" in w for w in result["warnings"])
    assert project["metadata"]["attached_proofs"]["github_proof"]["github_proof_id"] == proof["id"]


# ── Project Defense / Video finalization ──────────────────────────────────────


def test_defense_finalization_rejects_wrong_project_and_counts_pending(mem_store: dict) -> None:
    veribridge = _project(mem_store, title="VeriBridge")
    boston = _project(mem_store, title="Boston")
    session_id = str(uuid4())
    mem_store.setdefault("vbr_verification_sessions", {})[session_id] = {
        "id": session_id,
        "project_id": veribridge["id"],
        "status": "completed",
        "telemetry": {},
    }
    with pytest.raises(CanonicalEvidenceConflictError):
        finalize_proof_evidence(
            mem_store, user_id=USER_ID, proof_type="project_defense",
            proof_id=session_id, project_id=boston["id"],
        )
    result = finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="project_defense",
        proof_id=session_id, project_id=veribridge["id"],
    )
    # Unanalyzed defense: attached, pending, never counted.
    assert result["supported_skill_count"] == 0
    assert result["report_eligibility"] in ("project_attached", "skill_mapped")
    assert any("pending" in w.lower() or "not yet analyzed" in w.lower() for w in result["warnings"])
    links = list(mem_store.get("vbr_claim_evidence_links", {}).values())
    assert links and all(l["link_status"] == "pending" for l in links)


def test_video_finalization_is_citation_layer_never_counted(mem_store: dict) -> None:
    project = _project(mem_store, title="VeriBridge")
    video_id = str(uuid4())
    mem_store.setdefault("video_proofs", {})[video_id] = {
        "id": video_id,
        "user_id": USER_ID,
        "project_id": None,
        "claimed_skills": ["Python"],
    }
    result = finalize_proof_evidence(
        mem_store, user_id=USER_ID, proof_type="video",
        proof_id=video_id, project_id=project["id"],
    )
    assert result["supported_skill_count"] == 0
    links = list(mem_store.get("vbr_claim_evidence_links", {}).values())
    assert links and all(l["link_status"] == "excluded" for l in links)
    assert any("independent" in w.lower() for w in result["warnings"])
    assert mem_store["video_proofs"][video_id]["project_id"] == project["id"]


# ── Block-aware DOCX extraction ────────────────────────────────────────────────


def _rich_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Architecture", level=1)
    document.add_paragraph(
        "The backend uses FastAPI endpoints to serve the machine learning scoring API."
    )
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Technology"
    table.rows[1].cells[0].text = "REST API endpoints"
    table.rows[1].cells[1].text = "FastAPI request handling"
    document.add_heading("Results", level=1)
    document.add_paragraph(
        "Model accuracy reached 94.2% with an average latency of 120 ms per request."
    )
    code = document.add_paragraph(
        "def reroute(risk):\n    import fastapi\n    return risk"
    )
    code.style = document.styles["Normal"]
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_block_extraction_produces_typed_blocks_with_section_locators() -> None:
    extracted = extract_document_text(_rich_docx_bytes(), "rich.docx")
    assert extracted.status == "ok"
    kinds = {c.block_type for c in extracted.chunks}
    assert "heading" in kinds and "table" in kinds and "paragraph" in kinds
    assert "metric_result" in kinds  # "94.2% … accuracy/latency" paragraph
    assert "code_block" in kinds  # def/import/return heuristic
    table_chunk = next(c for c in extracted.chunks if c.block_type == "table")
    assert table_chunk.table_cells and table_chunk.table_cells[0][0] == "Component"
    # Paragraphs inherit the NEAREST REAL HEADING as their section locator —
    # never the layout style name "Normal".
    paragraph = next(
        c for c in extracted.chunks if c.block_type == "paragraph" and "FastAPI" in c.text
    )
    assert paragraph.section_label == "Architecture"


def test_docx_block_analysis_persists_block_locators_and_block_reasons() -> None:
    analysis = analyze_from_extracted_document(
        extract_document_text(_rich_docx_bytes(), "rich.docx")
    )
    table_evidence = [o for o in analysis.evidence_objects if o.get("block_type") == "table"]
    assert table_evidence, "expected the FastAPI table to produce block-typed evidence"
    assert table_evidence[0]["table_cells"]
    assert "table" in table_evidence[0]["reason"].lower()
    assert document_evidence_object_has_locator(table_evidence[0])
    # Section-anchored paragraphs carry a real heading, not "Normal".
    para_evidence = [
        o for o in analysis.evidence_objects if o.get("block_type") == "paragraph"
    ]
    assert all(o.get("section_label") != "Normal" for o in para_evidence)


def test_reextract_from_retained_original_is_additive_and_idempotent(mem_store: dict) -> None:
    doc_id = str(uuid4())
    original_objects = [
        {
            "skill_name": "Technical Documentation",
            "snippet": "It does not claim production benchmark results.",
            "reason": "Supported by document text mentioning results.",
            "section_label": "Normal",
        }
    ]
    mem_store.setdefault("optional_evidence_submissions", {})[doc_id] = {
        "id": doc_id,
        "user_id": USER_ID,
        "source_type": "document",
        "status": "analyzed",
        "file_path": "rich.docx",
        "analysis_json": {"title": "Rich doc"},
        "evidence_objects": list(original_objects),
    }
    artifact_id = str(uuid4())
    storage_path = f"{USER_ID}/document/{artifact_id}/rich.docx"
    mem_store.setdefault("proof_artifacts", {})[artifact_id] = {
        "id": artifact_id,
        "owner_user_id": USER_ID,
        "proof_type": "document",
        "artifact_type": "document_original",
        "proof_id": doc_id,
        "retained": True,
        "file_name": "rich.docx",
        "storage_path": storage_path,
    }
    mem_store.setdefault("_proof_artifact_objects", {})[storage_path] = _rich_docx_bytes()

    service = OptionalEvidenceService(mem_store)
    updated = service.reextract_blocks_from_retained_original(user_id=USER_ID, evidence_id=doc_id)
    assert updated is not None
    objects = updated["evidence_objects"]
    # Original object is preserved verbatim in position 0.
    assert objects[0]["snippet"] == original_objects[0]["snippet"]
    assert any(o.get("block_type") == "table" for o in objects)
    assert updated["analysis_json"]["block_extraction_version"] == 2

    count_after_first = len(objects)
    again = service.reextract_blocks_from_retained_original(user_id=USER_ID, evidence_id=doc_id)
    assert len(again["evidence_objects"]) == count_after_first  # idempotent


def test_reextract_is_owner_scoped(mem_store: dict) -> None:
    service = OptionalEvidenceService(mem_store)
    assert (
        service.reextract_blocks_from_retained_original(user_id=USER_ID, evidence_id=str(uuid4()))
        is None
    )

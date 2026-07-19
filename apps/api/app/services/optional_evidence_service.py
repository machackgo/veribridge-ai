"""Optional evidence boosters for Website Proof.

Text-first MVP for documents, profiles, and certificates.  It never scrapes
LinkedIn or invents metadata; every skill signal comes from submitted text.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any, Literal
from uuid import uuid4

SourceType = Literal["document", "linkedin_profile", "certificate_transcript"]

_TABLE = "optional_evidence_submissions"

_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
_MAX_EXTRACT_CHARS = 120_000  # ~30 k tokens — enough for analysis without memory blow-up


@dataclass
class DocumentChunk:
    """A piece of extracted text with optional provenance.

    ``block_type`` / ``block_index`` / ``table_cells`` / ``visual_description``
    / ``nearby_caption`` are the source-native block locators (migration-independent,
    persisted on each evidence object) that let reports cite "Page 2 · Table"
    instead of "document text mentions skill". They are extracted structurally
    (element kind + declared style + explicit caption references) — never
    inferred from prose content alone.
    """
    text: str
    page_number: int | None = None
    section_label: str | None = None
    line_start: int | None = None
    line_end: int | None = None
    block_type: str | None = None
    block_index: int | None = None
    table_cells: list[list[str]] | None = None
    visual_description: str | None = None
    nearby_caption: str | None = None
    figure_reference: str | None = None
    # Real adjacent prose (the paragraph immediately before a visual in the
    # same section) — extra matching context for caption-only visuals. Always
    # verbatim extracted text, never generated.
    context_text: str | None = None


@dataclass
class ExtractedDocument:
    file_name: str
    file_type: str  # pdf | docx | txt | md
    chunks: list[DocumentChunk] = field(default_factory=list)
    status: str = "ok"          # ok | no_text_found | unsupported_file | extraction_failed
    error_message: str | None = None

    @property
    def full_text(self) -> str:
        return "\n\n".join(c.text for c in self.chunks)

    @property
    def preview(self) -> str:
        return self.full_text[:600]


def extract_document_text(file_bytes: bytes, filename: str) -> ExtractedDocument:
    """Extract text + provenance from an uploaded document.

    Supported: PDF (page numbers), DOCX (paragraph index), TXT/MD (line ranges).
    Returns ExtractedDocument with status set accordingly — never raises.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(file_bytes, filename)
    if lower.endswith(".docx"):
        return _extract_docx(file_bytes, filename)
    if lower.endswith(".txt") or lower.endswith(".md"):
        return _extract_text(file_bytes, filename)

    ext = "." + filename.rsplit(".", 1)[-1] if "." in filename else ""
    return ExtractedDocument(
        file_name=filename,
        file_type=ext.lstrip(".") or "unknown",
        status="unsupported_file",
        error_message=f"File type '{ext}' is not supported. Upload PDF, DOCX, TXT, or MD.",
    )


def _extract_pdf(file_bytes: bytes, filename: str) -> ExtractedDocument:
    try:
        import pypdf  # optional dep — checked at call time
    except ImportError:
        return ExtractedDocument(
            file_name=filename, file_type="pdf",
            status="extraction_failed",
            error_message="PDF extraction library (pypdf) is not installed on this server.",
        )
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        chunks: list[DocumentChunk] = []
        total_chars = 0
        for i, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception:
                page_text = ""
            page_text = page_text.strip()
            if not page_text:
                continue
            chunks.append(DocumentChunk(text=page_text, page_number=i))
            total_chars += len(page_text)
            if total_chars >= _MAX_EXTRACT_CHARS:
                break
        if not chunks:
            return ExtractedDocument(
                file_name=filename, file_type="pdf",
                status="no_text_found",
                error_message="No selectable text found. OCR document extraction not yet available.",
            )
        return ExtractedDocument(file_name=filename, file_type="pdf", chunks=chunks)
    except Exception as exc:
        return ExtractedDocument(
            file_name=filename, file_type="pdf",
            status="extraction_failed",
            error_message=f"PDF extraction error: {exc}",
        )


# Paragraph styles that are layout defaults, not real section locators. A skill
# match inside a bare default-styled paragraph is "text mentions skill" — it must
# never masquerade as a section citation.
_NON_LOCATOR_STYLES = frozenset({"normal", "body text", "default", "default paragraph font"})

# Caption keyword → canonical visual block type (closed vocabulary; anything
# ambiguous stays the generic type and the report fails closed downstream).
_CAPTION_BLOCK_TYPES = (
    ("architecture", "architecture_diagram"),
    ("flowchart", "architecture_diagram"),
    ("diagram", "architecture_diagram"),
    ("screenshot", "screenshot"),
    ("ui ", "screenshot"),
    ("interface", "screenshot"),
    ("bar chart", "chart"),
    ("line chart", "chart"),
    ("chart", "chart"),
    ("graph", "graph"),
)

_FIGURE_REF_RE = re.compile(r"\b((?:figure|table|chart|diagram|exhibit)\s*\d+)\b", re.I)
_CODE_TOKEN_RE = re.compile(
    r"(\bdef\s+\w+\(|\bimport\s+\w+|\breturn\b|[{};]\s*$|=>|::|\bclass\s+\w+[(:]|\bawait\s|\basync\s)",
    re.M,
)
# One SOURCE LINE of code split across single-line paragraphs (a very common
# way rich documents embed code). Conservative, structure-only signals.
_CODE_LINE_RE = re.compile(
    r"(^@\w+[.(]|^def\s+\w+\(|^class\s+\w+[(:]|^return\b|^from\s+\w+\s+import\b|"
    r"^import\s+\w+|\w+\s*=\s*\w+[\w.]*\(.*\)$|:\s*$|\)\s*:$|^\}|^\{)",
)
_METRIC_RE = re.compile(r"\b\d+(?:\.\d+)?\s*(?:%|ms\b|s\b|mb\b|gb\b|rps\b|qps\b)", re.I)
_METRIC_TERMS = ("accuracy", "latency", "throughput", "precision", "recall", "uptime", "coverage", "f1", "error rate", "response time")


def _classify_visual_block(caption: str, *, is_chart_part: bool) -> str:
    lower = (caption or "").lower()
    for keyword, block_type in _CAPTION_BLOCK_TYPES:
        if keyword in lower:
            return block_type
    return "chart" if is_chart_part else "image"


def _looks_like_code(text: str, style_name: str) -> bool:
    if any(token in style_name for token in ("code", "preformatted", "macro")):
        return True
    return len(_CODE_TOKEN_RE.findall(text)) >= 2


def _looks_like_metric(text: str) -> bool:
    lower = text.lower()
    return bool(_METRIC_RE.search(text)) and any(term in lower for term in _METRIC_TERMS)


def _docx_table_cells(table: Any) -> list[list[str]]:
    """Bounded, safe cell text for one table (≤8 rows × ≤6 cols, ≤80 chars each)."""
    cells: list[list[str]] = []
    for row in table.rows[:8]:
        out_row = [cell.text.strip()[:80] for cell in row.cells[:6]]
        if any(out_row):
            cells.append(out_row)
    return cells


def _extract_docx(file_bytes: bytes, filename: str) -> ExtractedDocument:
    """Block-aware DOCX extraction.

    Walks the document body IN ORDER (paragraphs and tables), classifying each
    real block: headings become section locators for the paragraphs beneath
    them, tables keep their cell structure, embedded drawings/charts become
    visual blocks described by their own caption/alt-text, and code/metric
    paragraphs are typed structurally. No page numbers are invented — DOCX has
    no fixed pagination, so locators are section + block index, stated as such.
    """
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return ExtractedDocument(
            file_name=filename, file_type="docx",
            status="extraction_failed",
            error_message="DOCX extraction library (python-docx) is not installed on this server.",
        )
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        chunks: list[DocumentChunk] = []
        total_chars = 0
        current_heading: str | None = None
        block_index = 0
        pending_visuals: list[DocumentChunk] = []
        # The last real prose paragraph in the current section — verbatim
        # adjacent context for a caption-only visual (never generated text).
        last_prose: str | None = None
        # Consecutive single-line code paragraphs buffered into one code block.
        code_run: list[str] = []
        code_run_start_index = 0

        def _flush_code_run() -> None:
            nonlocal code_run
            if not code_run:
                return
            if len(code_run) >= 2:
                chunks.append(
                    DocumentChunk(
                        text="\n".join(code_run)[:1200],
                        section_label=current_heading,
                        block_type="code_block",
                        block_index=code_run_start_index,
                    )
                )
            else:
                chunks.append(
                    DocumentChunk(
                        text=code_run[0],
                        section_label=current_heading,
                        block_type="paragraph",
                        block_index=code_run_start_index,
                    )
                )
            code_run = []

        def _flush_pending_caption(caption: str) -> None:
            """Attach a Figure/Table-style caption to the visual block just above."""
            for visual in pending_visuals:
                visual.nearby_caption = caption[:200]
                ref = _FIGURE_REF_RE.search(caption)
                if ref:
                    visual.figure_reference = ref.group(1).title()
                visual.block_type = _classify_visual_block(
                    caption, is_chart_part=visual.block_type == "chart"
                )
                visual.text = caption.strip()[:280]
                visual.visual_description = (
                    f"Embedded visual ({visual.block_type.replace('_', ' ')}) captioned: "
                    f"{caption.strip()[:160]}"
                )
            pending_visuals.clear()

        # python-docx: iterate body elements in document order.
        body = doc.element.body
        for element in body.iterchildren():
            if total_chars >= _MAX_EXTRACT_CHARS:
                break
            tag = element.tag.rsplit("}", 1)[-1]
            if tag == "tbl":
                _flush_code_run()
                block_index += 1
                table = Table(element, doc)
                cells = _docx_table_cells(table)
                if not cells:
                    continue
                text = "\n".join(" | ".join(c for c in row if c) for row in cells)
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="table",
                        block_index=block_index,
                        table_cells=cells,
                    )
                )
                pending_visuals.clear()
                total_chars += len(text)
                continue
            if tag != "p":
                continue
            block_index += 1
            para = Paragraph(element, doc)
            text = para.text.strip()
            style_name = (para.style.name if para.style else "") or ""
            style_lower = style_name.lower()

            # Embedded drawings (images / native charts) — one visual block per
            # drawing. Alt text (docPr descr/name) is the only content read.
            drawings = element.findall(
                ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
            )
            if drawings:
                _flush_code_run()
            for drawing in drawings:
                doc_pr = drawing.find(
                    ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
                )
                alt = ""
                if doc_pr is not None:
                    alt = (doc_pr.get("descr") or doc_pr.get("name") or "").strip()
                is_chart = bool(
                    drawing.findall(
                        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphicData"
                        "[@uri='http://schemas.openxmlformats.org/drawingml/2006/chart']"
                    )
                )
                visual = DocumentChunk(
                    text=alt[:280],
                    section_label=current_heading,
                    block_type="chart" if is_chart else _classify_visual_block(alt, is_chart_part=is_chart),
                    block_index=block_index,
                    visual_description=(
                        f"Embedded {'chart' if is_chart else 'image'}"
                        + (f" (alt text: {alt[:120]})" if alt else " with no alt text")
                    ),
                    # Verbatim adjacent prose in the same section — matching
                    # context for a caption-only visual.
                    context_text=(last_prose or "")[:300] or None,
                )
                chunks.append(visual)
                pending_visuals.append(visual)

            if not text:
                continue

            is_caption = "caption" in style_lower or bool(_FIGURE_REF_RE.match(text))
            # A single SHORT line of code (documents often split code across
            # one-line paragraphs) — buffered into a merged code block. A
            # paragraph that is already strongly code-typed on its own
            # (≥2 structural tokens) goes straight to the code_block branch.
            is_code_line = (
                bool(_CODE_LINE_RE.search(text))
                and not is_caption
                and len(text) <= 160
                and not _looks_like_code(text, style_lower)
            )

            if not is_code_line:
                _flush_code_run()

            if style_lower.startswith(("heading", "title", "subtitle")):
                current_heading = text[:120]
                last_prose = None
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="heading",
                        block_index=block_index,
                    )
                )
                pending_visuals.clear()
            elif is_caption:
                if pending_visuals:
                    _flush_pending_caption(text)
                else:
                    chunks.append(
                        DocumentChunk(
                            text=text,
                            section_label=current_heading,
                            block_type="caption",
                            block_index=block_index,
                            figure_reference=(
                                _FIGURE_REF_RE.search(text).group(1).title()
                                if _FIGURE_REF_RE.search(text)
                                else None
                            ),
                        )
                    )
            elif is_code_line:
                if not code_run:
                    code_run_start_index = block_index
                code_run.append(text)
            elif _looks_like_code(text, style_lower):
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="code_block",
                        block_index=block_index,
                    )
                )
                pending_visuals.clear()
            elif _looks_like_metric(text):
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="metric_result",
                        block_index=block_index,
                    )
                )
                pending_visuals.clear()
                last_prose = text
            elif style_lower.startswith("list") or style_lower.startswith("bullet"):
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="list",
                        block_index=block_index,
                    )
                )
                pending_visuals.clear()
            else:
                chunks.append(
                    DocumentChunk(
                        text=text,
                        section_label=current_heading,
                        block_type="paragraph",
                        block_index=block_index,
                    )
                )
                pending_visuals.clear()
                last_prose = text
            total_chars += len(text)
        _flush_code_run()

        # A code block's meaning is often stated by the paragraph right after it
        # ("Evidence interpretation: this code block supports …"). Attach that
        # VERBATIM adjacent paragraph as matching context — deterministic
        # adjacency, only when the paragraph explicitly references the code.
        for position, chunk in enumerate(chunks):
            if chunk.block_type != "code_block" or chunk.context_text:
                continue
            for follower in chunks[position + 1 : position + 3]:
                if follower.block_type == "paragraph" and re.search(
                    r"\b(this code|code block|snippet)\b", follower.text, re.I
                ):
                    chunk.context_text = follower.text[:300]
                    break

        # Visual blocks with no caption and no alt text carry no text to match —
        # keep them (they are honest structure) but they can never create a
        # skill claim by themselves.
        chunks = [c for c in chunks if c.text or c.block_type in ("image", "chart", "graph", "screenshot", "architecture_diagram", "table")]
        if not any(c.text for c in chunks):
            return ExtractedDocument(
                file_name=filename, file_type="docx",
                status="no_text_found",
                error_message="No text paragraphs found in this DOCX file.",
            )
        return ExtractedDocument(file_name=filename, file_type="docx", chunks=chunks)
    except Exception as exc:
        return ExtractedDocument(
            file_name=filename, file_type="docx",
            status="extraction_failed",
            error_message=f"DOCX extraction error: {exc}",
        )


def _extract_text(file_bytes: bytes, filename: str) -> ExtractedDocument:
    file_type = "md" if filename.lower().endswith(".md") else "txt"
    try:
        raw = file_bytes.decode("utf-8", errors="replace")
        lines = raw.splitlines()
        chunks: list[DocumentChunk] = []
        total_chars = 0
        para_lines: list[str] = []
        para_start: int | None = None
        for i, line in enumerate(lines, start=1):
            if line.strip():
                if para_start is None:
                    para_start = i
                para_lines.append(line)
            else:
                if para_lines:
                    text = " ".join(para_lines).strip()
                    chunks.append(DocumentChunk(
                        text=text, line_start=para_start, line_end=i - 1,
                    ))
                    total_chars += len(text)
                    para_lines = []
                    para_start = None
                    if total_chars >= _MAX_EXTRACT_CHARS:
                        break
        if para_lines:
            text = " ".join(para_lines).strip()
            chunks.append(DocumentChunk(text=text, line_start=para_start, line_end=len(lines)))
        if not chunks:
            return ExtractedDocument(
                file_name=filename, file_type=file_type,
                status="no_text_found",
                error_message="No text content found in this file.",
            )
        return ExtractedDocument(file_name=filename, file_type=file_type, chunks=chunks)
    except Exception as exc:
        return ExtractedDocument(
            file_name=filename, file_type=file_type,
            status="extraction_failed",
            error_message=f"Text extraction error: {exc}",
        )

@dataclass(frozen=True)
class SkillRule:
    skill: str
    phrases: tuple[str, ...]
    context_terms: tuple[str, ...] = ()
    tool_terms: tuple[str, ...] = ()
    task_terms: tuple[str, ...] = ()
    result_terms: tuple[str, ...] = ()
    exclude_terms: tuple[str, ...] = ()


_GENERAL_TASK_TERMS = (
    "built", "implemented", "trained", "fine-tuned", "evaluated", "optimized",
    "deployed", "processed", "generated", "classified", "detected", "render",
    "renders", "rendered", "queried", "stored", "validated", "monitored",
)
_PROJECT_CONTEXT_TERMS = (
    "project", "application", "app", "system", "pipeline", "workflow", "model",
    "api", "dashboard", "frontend", "backend", "browser", "dataset", "results",
    "accuracy", "latency", "scene", "mesh", "geometry", "report",
)


_SKILL_RULES: tuple[SkillRule, ...] = (
    # 3D graphics / browser rendering
    SkillRule("Three.js", ("three.js", "three js", "threejs"), tool_terms=("three.js", "three js", "threejs"), task_terms=("render", "renders", "rendered", "implemented", "built"), result_terms=("interactive", "3d", "webgl", "scene", "object")),
    SkillRule("WebGL", ("webgl", "web gl"), tool_terms=("webgl", "web gl"), task_terms=("render", "renders", "rendered", "implemented"), result_terms=("browser", "3d", "graphics", "scene")),
    SkillRule("Interactive 3D Graphics", ("3d graphics", "interactive 3d", "3d rendering", "3d scene", "3d object"), context_terms=("render", "graphics", "scene", "object", "webgl", "three.js", "browser")),
    SkillRule("Computer Graphics", ("3d graphics", "interactive 3d", "3d rendering", "computer graphics", "geometry processing", "mesh processing", "scene", "camera", "controls", "material", "geometry"), context_terms=("3d", "webgl", "three.js", "rendering", "graphics", "mesh", "geometry", "browser")),
    SkillRule("Geometry Optimization", ("geometry optimization", "geometry processing", "mesh processing", "mesh simplification", "geometry simplifier"), task_terms=("optimized", "optimization", "simplification", "processing", "simplifier"), result_terms=("geometry", "mesh", "performance")),
    SkillRule("3D Mesh Simplification", ("mesh simplification", "simplifier", "geometry simplifier", "mesh processing"), context_terms=("mesh", "geometry", "3d", "graphics")),
    SkillRule("Frontend Development", ("browser rendering", "rendering pipeline", "javascript modules", "browser interaction", "frontend application", "frontend development", "react", "next.js", "nextjs", "typescript", "javascript"), context_terms=("frontend", "browser", "ui", "client", "react", "next", "javascript", "typescript", "3d", "webgl")),
    SkillRule("Browser Rendering", ("browser rendering", "rendering pipeline", "webgl rendering", "browser interaction"), context_terms=("browser", "webgl", "frontend", "rendering")),
    # AI / ML / data science
    SkillRule("Machine Learning", ("machine learning", "model training", "supervised learning", "classification model", "regression model", "cnn model"), task_terms=("trained", "evaluated", "classified", "predicted"), result_terms=("accuracy", "f1", "precision", "recall", "dataset")),
    SkillRule("Deep Learning", ("deep learning", "neural network", "cnn", "convolutional", "transformer", "pytorch", "tensorflow"), context_terms=("model", "training", "classification", "vision", "nlp", "dataset")),
    SkillRule("TensorFlow", ("tensorflow", "tf.keras", "keras"), tool_terms=("tensorflow", "tf.keras", "keras"), task_terms=("trained", "evaluated", "implemented"), result_terms=("model", "accuracy", "f1", "dataset")),
    SkillRule("NLP", ("nlp", "natural language processing", "text classification", "named entity recognition", "sentiment analysis", "llm assistant"), context_terms=("text", "language", "tokens", "embedding", "llm", "rag", "retrieval", "documents")),
    SkillRule("LLM", ("llm", "large language model", "gpt", "claude", "llama", "prompt engineering"), context_terms=("prompt", "generation", "rag", "retrieval", "embedding", "model")),
    SkillRule("RAG", ("rag", "retrieval augmented generation", "retrieval-augmented generation"), context_terms=("retrieval", "embedding", "vector", "documents", "llm")),
    SkillRule("Embeddings", ("embedding", "embeddings", "vector search", "semantic search", "vector database"), context_terms=("vector", "semantic", "retrieval", "rag", "documents")),
    SkillRule("Computer Vision", ("computer vision", "image recognition", "image classification", "opencv", "cnn for images", "object detection", "segmentation", "detection model", "webcam", "image processing", "visual perception", "pose detection", "ocr"), context_terms=("image", "video", "vision", "detection", "classification", "ocr", "webcam", "segmentation", "pose"), exclude_terms=("data visualization", "visualization", "visual analytics", "visual output", "svg", "chart", "charts", "graph gallery")),
    SkillRule("Object Detection", ("object detection", "yolo", "bounding box", "bounding boxes"), context_terms=("image", "vision", "detected", "model")),
    SkillRule("OCR", ("ocr", "text extraction", "optical character recognition"), context_terms=("image", "document", "text", "extraction")),
    SkillRule("Model Evaluation", ("accuracy", "f1-score", "f1 score", "roc-auc", "roc auc", "precision", "recall", "confusion matrix"), context_terms=("model", "evaluation", "evaluated", "test set", "validation")),
    SkillRule("MLOps", ("mlops", "model monitoring", "model registry", "model deployment", "experiment tracking", "drift monitoring"), context_terms=("model", "monitoring", "deployment", "tracking", "pipeline")),
    SkillRule("Data Science", ("data science", "data scientist", "data analysis", "recommendation system", "feature engineering"), context_terms=("dataset", "analysis", "model", "pandas", "visualization", "recommendation")),
    SkillRule("Data Analysis", ("data analysis", "exploratory analysis", "eda", "pandas", "numpy"), context_terms=("dataset", "analysis", "cleaning", "features")),
    SkillRule("Data Visualization", ("data visualization", "visual analytics", "visual output", "dashboard", "chart", "charts", "plotly", "d3.js", "d3 ", "matplotlib", "seaborn", "svg visualization"), context_terms=("dashboard", "chart", "charts", "visualization", "results", "analytics", "svg", "graph")),
    SkillRule("D3.js", ("d3.js", "d3 ", "d3 graph", "d3 chart"), tool_terms=("d3.js", "d3 "), task_terms=("built", "implemented", "rendered", "visualized"), result_terms=("chart", "charts", "svg", "graph", "visualization", "interactive")),
    SkillRule("Interactive Charts", ("interactive chart", "interactive charts", "chart interaction", "chart interactions", "interactive visualization", "hover tooltip", "tooltip", "brush", "zoom"), context_terms=("chart", "charts", "visualization", "svg", "d3", "interaction", "tooltip")),
    SkillRule("SVG Visualization", ("svg visualization", "svg chart", "svg charts", "svg", "scalable vector graphics"), context_terms=("visualization", "chart", "charts", "graph", "d3", "frontend")),
    SkillRule("Graph Analysis", ("graph analysis", "graph analytics", "graph gallery", "network graph", "graph visualization", "node-link", "nodes and edges", "nodes", "edges"), context_terms=("graph", "chart", "visualization", "network", "nodes", "edges", "analytics")),
    SkillRule("Python", ("python", "pandas", "numpy", "scikit-learn", "sklearn"), tool_terms=("python", "pandas", "numpy", "scikit-learn", "sklearn"), task_terms=("implemented", "trained", "analyzed", "processed")),
    # Backend / frontend / databases
    SkillRule("Backend API", ("rest api", "backend api", "api endpoint", "fastapi endpoint", "server endpoint"), context_terms=("api", "endpoint", "backend", "server", "request", "response")),
    SkillRule("API Development", ("api development", "rest api", "endpoint", "http api", "openapi", "swagger"), context_terms=("api", "endpoint", "http", "request", "response")),
    SkillRule("FastAPI", ("fastapi", "fast api"), tool_terms=("fastapi", "fast api"), task_terms=("implemented", "built", "served", "validated"), result_terms=("api", "endpoint", "backend")),
    SkillRule("Flask", ("flask", "flask api"), tool_terms=("flask",), task_terms=("implemented", "built", "served"), result_terms=("api", "backend", "endpoint")),
    SkillRule("Django", ("django", "django rest framework", "drf"), tool_terms=("django", "django rest framework", "drf"), task_terms=("implemented", "built", "served"), result_terms=("api", "backend", "endpoint")),
    SkillRule("React", ("react", "react.js", "reactjs"), tool_terms=("react", "react.js", "reactjs"), task_terms=("built", "implemented", "rendered"), result_terms=("ui", "component", "frontend", "application")),
    SkillRule("Next.js", ("next.js", "nextjs", "next js"), tool_terms=("next.js", "nextjs", "next js"), task_terms=("built", "implemented", "rendered"), result_terms=("frontend", "application", "route", "server")),
    SkillRule("JavaScript", ("javascript", "javascript modules", "node.js", "d3.js", "d3 "), tool_terms=("javascript", "node.js", "d3.js"), task_terms=("implemented", "built", "rendered"), result_terms=("frontend", "browser", "module", "application")),
    SkillRule("TypeScript", ("typescript", "tsx", "ts-node"), tool_terms=("typescript", "tsx"), task_terms=("implemented", "typed", "built"), result_terms=("frontend", "api", "application")),
    SkillRule("Database", ("database", "sql", "nosql", "schema", "query", "queries"), context_terms=("stored", "queried", "tables", "records", "postgres", "mongodb", "supabase")),
    SkillRule("SQL", ("sql", "sql query", "sql database"), context_terms=("query", "database", "table", "postgres", "schema")),
    SkillRule("PostgreSQL", ("postgresql", "postgres"), tool_terms=("postgresql", "postgres"), task_terms=("queried", "stored", "joined", "indexed"), result_terms=("database", "sql", "schema")),
    SkillRule("MongoDB", ("mongodb", "mongo db", "mongoose"), tool_terms=("mongodb", "mongo db", "mongoose"), task_terms=("queried", "stored", "modeled"), result_terms=("database", "collection", "document")),
    SkillRule("Supabase", ("supabase", "row level security", "rls"), tool_terms=("supabase",), task_terms=("implemented", "queried", "stored"), result_terms=("database", "auth", "postgres", "storage")),
    # Deployment / cloud. Intentionally excludes generic "pipeline".
    SkillRule("Docker", ("docker", "container", "dockerfile"), tool_terms=("docker", "dockerfile"), task_terms=("containerized", "deployed", "built"), result_terms=("container", "image", "service")),
    SkillRule("DevOps", ("docker", "ci/cd", "kubernetes", "deployment pipeline", "cloud deployment", "monitoring", "infrastructure", "github actions", "container", "server", "nginx", "terraform"), context_terms=("deploy", "deployment", "infrastructure", "container", "cloud", "ci/cd", "monitoring", "server", "nginx", "terraform", "github actions")),
    SkillRule("Cloud Deployment", ("cloud deployment", "aws", "gcp", "azure", "vercel", "render.com", "railway"), context_terms=("deploy", "deployment", "cloud", "hosting", "service")),
    SkillRule("CI/CD", ("ci/cd", "github actions", "gitlab ci", "deployment pipeline"), context_terms=("build", "test", "deploy", "workflow", "pipeline")),
    SkillRule("Technical Documentation", ("technical documentation", "project report", "methodology", "results", "limitations"), context_terms=("report", "methodology", "results", "limitations", "documentation", "project")),
)

_ISSUER_PATTERNS = (
    r"(?:issuer|issued by|institution|university|provider)[:\s]+([A-Z][A-Za-z0-9 &.,-]{2,80})",
    r"(Coursera|edX|Udacity|Stanford University|MIT|Google|Microsoft|AWS|Amazon Web Services|IBM)",
)
_TITLE_PATTERNS = (
    r"(?:certificate|course|program|transcript)[:\s]+([A-Z][A-Za-z0-9 &/.,:+-]{3,100})",
    r"(Machine Learning|Deep Learning|Data Science|Cloud Practitioner|Computer Vision|Full Stack|Backend API)",
)
_DATE_PATTERN = r"\b(?:20\d{2}|19\d{2})(?:[-/]\d{1,2})?(?:[-/]\d{1,2})?\b"


@dataclass
class OptionalEvidenceAnalysis:
    status: str
    source_type: SourceType
    summary: str
    evidence_objects: list[dict[str, Any]]
    analysis_json: dict[str, Any]


def _sentences(text: str) -> list[str]:
    chunks = re.split(r"(?<=[.!?])\s+|\n+", text or "")
    return [c.strip() for c in chunks if len(c.strip()) >= 12]


def _page_chunks(text: str) -> list[tuple[int | None, str]]:
    # Accept explicit markers from callers/tests: "\f" or "--- page 2 ---".
    if "\f" in text:
        return [(i + 1, part.strip()) for i, part in enumerate(text.split("\f")) if part.strip()]
    marker = re.compile(r"---\s*page\s+(\d+)\s*---", re.I)
    matches = list(marker.finditer(text))
    if not matches:
        return [(None, text)]
    chunks: list[tuple[int | None, str]] = []
    for idx, match in enumerate(matches):
        start = match.end()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        chunks.append((int(match.group(1)), text[start:end].strip()))
    return [(p, c) for p, c in chunks if c]


def _matched_phrases(rule: SkillRule, lower: str) -> list[str]:
    if rule.exclude_terms and any(term in lower for term in rule.exclude_terms):
        return []
    hits = [phrase for phrase in rule.phrases if phrase in lower]
    if not hits:
        return []
    if rule.context_terms and not any(term in lower for term in rule.context_terms):
        return []
    return hits


def _confidence(snippet: str, source_type: SourceType, rule: SkillRule | None = None) -> str:
    lower = snippet.lower()
    tool_terms = rule.tool_terms if rule and rule.tool_terms else rule.phrases if rule else ()
    task_terms = rule.task_terms if rule and rule.task_terms else _GENERAL_TASK_TERMS
    result_terms = rule.result_terms if rule and rule.result_terms else _PROJECT_CONTEXT_TERMS
    has_tool = bool(tool_terms) and any(term in lower for term in tool_terms)
    has_task = any(term in lower for term in task_terms)
    has_result = any(term in lower for term in result_terms)
    detail = sum(1 for kw in _GENERAL_TASK_TERMS + _PROJECT_CONTEXT_TERMS if kw in lower)
    if source_type == "certificate_transcript":
        return "medium"
    if has_tool and has_task and has_result:
        return "high"
    if detail >= 2:
        return "high"
    if has_tool or len(snippet.split()) >= 8:
        return "medium"
    return "low"


def _find_meta(text: str) -> dict[str, str | None]:
    def first(patterns: tuple[str, ...]) -> str | None:
        for pat in patterns:
            m = re.search(pat, text or "", re.I)
            if m:
                return m.group(1).strip(" .,-")
        return None

    date_match = re.search(_DATE_PATTERN, text or "")
    return {
        "issuer": first(_ISSUER_PATTERNS),
        "title": first(_TITLE_PATTERNS),
        "date": date_match.group(0) if date_match else None,
    }


# Block types whose FULL bounded text is a matching candidate (tables/captions/
# code/metric rows are structural units, not prose sentences).
_WHOLE_BLOCK_TYPES = frozenset(
    {"table", "chart", "graph", "image", "screenshot", "architecture_diagram",
     "code_block", "metric_result", "caption", "list"}
)

_MAX_EVIDENCE_OBJECTS = 40

_BLOCK_REASON_LABEL = {
    "table": "a table",
    "chart": "a chart",
    "graph": "a graph",
    "image": "an embedded image",
    "screenshot": "a UI screenshot",
    "architecture_diagram": "an architecture diagram",
    "code_block": "a code block",
    "metric_result": "a reported metric/result",
    "caption": "a figure/table caption",
    "list": "a list",
    "heading": "a section heading",
}


def _block_reason(block: dict[str, Any] | None, source_type: SourceType, hits: list[str]) -> str:
    mentioned = ", ".join(hits[:3])
    block_type = str((block or {}).get("block_type") or "")
    label = _BLOCK_REASON_LABEL.get(block_type)
    if label:
        where = str((block or {}).get("section_label") or "").strip()
        suffix = f" in the '{where}' section" if where else ""
        return f"Supported by {label}{suffix} referencing {mentioned}."
    return f"Supported by {source_type.replace('_', ' ')} text mentioning {mentioned}."


def _analyze_chunks(
    chunks: list[tuple[int | None, str | None, str]],
    *,
    source_type: SourceType,
    profile_url: str | None,
    section_label: str | None,
    file_path: str | None,
    file_name: str | None,
    file_type: str | None,
    meta: dict[str, Any],
    line_refs: list[tuple[int | None, int | None]] | None = None,
    block_meta: list[dict[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], str, str]:
    """Core extraction loop; chunks = [(page_num, section_label, text)].

    ``block_meta`` (aligned with ``chunks``) carries the source-native block
    locators from block-aware extraction; each evidence object persists them so
    reports can cite the exact table/chart/diagram/code block, not just "text".
    """
    evidence: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    evidence_type = {
        "document": "document_snippet",
        "linkedin_profile": "profile_snippet",
        "certificate_transcript": "certificate_or_transcript_snippet",
    }[source_type]

    for chunk_idx, (page_number, chunk_section, chunk_text) in enumerate(chunks):
        block = block_meta[chunk_idx] if block_meta and chunk_idx < len(block_meta) else None
        block_type = str((block or {}).get("block_type") or "") or None
        candidates = _sentences(chunk_text)
        # Structural blocks match on their whole bounded content (a table row set
        # or caption rarely splits into prose sentences).
        if block_type in _WHOLE_BLOCK_TYPES:
            whole = (chunk_text or "").strip()
            caption = str((block or {}).get("nearby_caption") or "").strip()
            # Verbatim adjacent prose (extracted, never generated) participates
            # in matching for caption-only visuals; the caption stays the cited
            # excerpt. The whole block is ONE matching unit — a table matched by
            # a skill yields one block citation, never one object per row (no
            # raw evidence dumps, no inflated counts).
            context = str((block or {}).get("context_text") or "").strip()
            merged = " ".join(part for part in (whole, caption, context) if part)
            candidates = [merged[:600]] if merged else []
        for sentence in candidates:
            lower = sentence.lower()
            for rule in _SKILL_RULES:
                hits = _matched_phrases(rule, lower)
                if not hits:
                    continue
                skill = rule.skill
                key = (skill, sentence[:120])
                if key in seen:
                    continue
                seen.add(key)
                conf = _confidence(sentence, source_type, rule)
                lr = line_refs[chunk_idx] if line_refs and chunk_idx < len(line_refs) else (None, None)
                evidence.append({
                    "source_type": source_type,
                    "evidence_type": (
                        f"document_{block_type}" if block_type and block_type != "paragraph" else evidence_type
                    ),
                    "skill_name": skill,
                    "confidence": conf,
                    "snippet": sentence[:280],
                    "reason": _block_reason(block, source_type, hits),
                    "page_number": page_number,
                    "profile_url": profile_url,
                    "section_label": chunk_section or section_label,
                    "file_path": file_path,
                    "file_name": file_name,
                    "file_type": file_type,
                    "line_start": lr[0],
                    "line_end": lr[1],
                    "issuer": meta.get("issuer"),
                    "title": meta.get("title"),
                    "date": meta.get("date"),
                    # Source-native block locators (None on non-block extraction).
                    "block_type": block_type,
                    "block_index": (block or {}).get("block_index"),
                    "table_cells": (block or {}).get("table_cells") or [],
                    "visual_description": (block or {}).get("visual_description"),
                    "nearby_caption": (block or {}).get("nearby_caption"),
                    "figure_reference": (block or {}).get("figure_reference"),
                })
            if len(evidence) >= _MAX_EVIDENCE_OBJECTS:
                break
        if len(evidence) >= _MAX_EVIDENCE_OBJECTS:
            break

    status = "analyzed" if evidence else "needs_review"
    summary = (
        f"Extracted {len(evidence)} evidence item(s)."
        if evidence else "No concrete skill evidence found in submitted text."
    )
    return evidence, status, summary


def analyze_optional_evidence(
    *,
    source_type: SourceType,
    raw_text: str = "",
    profile_url: str | None = None,
    section_label: str | None = None,
    file_path: str | None = None,
    file_name: str | None = None,
    file_type: str | None = None,
) -> OptionalEvidenceAnalysis:
    text = (raw_text or "").strip()
    if not text and source_type == "linkedin_profile" and profile_url:
        return OptionalEvidenceAnalysis(
            status="url_added",
            source_type=source_type,
            summary="Profile URL added. Paste public profile text for skill analysis.",
            evidence_objects=[],
            analysis_json={"profile_url": profile_url, "needs_more_detail": True},
        )
    if len(text) < 20:
        return OptionalEvidenceAnalysis(
            status="needs_review",
            source_type=source_type,
            summary="Evidence text is too short to extract reliable skills.",
            evidence_objects=[],
            analysis_json={"profile_url": profile_url, "file_path": file_path},
        )

    meta = _find_meta(text) if source_type == "certificate_transcript" else {}
    raw_chunks = [(p, None, t) for p, t in _page_chunks(text)]
    evidence, status, summary = _analyze_chunks(
        raw_chunks, source_type=source_type, profile_url=profile_url,
        section_label=section_label, file_path=file_path,
        file_name=file_name, file_type=file_type, meta=meta,
    )
    return OptionalEvidenceAnalysis(
        status=status,
        source_type=source_type,
        summary=summary,
        evidence_objects=evidence,
        analysis_json={
            "source_type": source_type,
            "summary": summary,
            "detected_skills": sorted({e["skill_name"] for e in evidence}),
            "profile_url": profile_url,
            "section_label": section_label,
            "file_path": file_path,
            "file_name": file_name,
            "file_type": file_type,
            **meta,
        },
    )


def analyze_from_extracted_document(
    doc: ExtractedDocument,
    *,
    source_type: SourceType = "document",
    profile_url: str | None = None,
) -> OptionalEvidenceAnalysis:
    """Analyze an already-extracted document's chunks with full provenance."""
    if doc.status != "ok":
        return OptionalEvidenceAnalysis(
            status=doc.status,
            source_type=source_type,
            summary=doc.error_message or "Document extraction failed.",
            evidence_objects=[],
            analysis_json={
                "source_type": source_type,
                "file_name": doc.file_name,
                "file_type": doc.file_type,
                "extraction_status": doc.status,
                "error": doc.error_message,
            },
        )

    raw_chunks = [
        (c.page_number, c.section_label, c.text) for c in doc.chunks
    ]
    line_refs = [(c.line_start, c.line_end) for c in doc.chunks]
    block_meta = [
        {
            "block_type": c.block_type,
            "block_index": c.block_index,
            "table_cells": c.table_cells,
            "visual_description": c.visual_description,
            "nearby_caption": c.nearby_caption,
            "figure_reference": c.figure_reference,
            "section_label": c.section_label,
            "context_text": c.context_text,
        }
        for c in doc.chunks
    ]
    meta = _find_meta(doc.full_text) if source_type == "certificate_transcript" else {}

    evidence, status, summary = _analyze_chunks(
        raw_chunks, source_type=source_type, profile_url=profile_url,
        section_label=None, file_path=None,
        file_name=doc.file_name, file_type=doc.file_type,
        meta=meta, line_refs=line_refs, block_meta=block_meta,
    )
    return OptionalEvidenceAnalysis(
        status=status,
        source_type=source_type,
        summary=summary,
        evidence_objects=evidence,
        analysis_json={
            "source_type": source_type,
            "summary": summary,
            "detected_skills": sorted({e["skill_name"] for e in evidence}),
            "file_name": doc.file_name,
            "file_type": doc.file_type,
            "extracted_text_preview": doc.preview,
            **meta,
        },
    )


class OptionalEvidencePersistError(RuntimeError):
    """Raised when an optional evidence submission could not be persisted."""


class OptionalEvidenceService:
    def __init__(self, db: Any) -> None:
        self._db = db

    def _persist(self, payload: dict[str, Any], *, strict: bool = False) -> dict[str, Any]:
        if isinstance(self._db, dict):
            now = datetime.now(UTC).isoformat()
            row = {"id": str(uuid4()), "created_at": now, "updated_at": now, **payload}
            self._db.setdefault(_TABLE, {})[row["id"]] = row
            return row
        try:
            resp = self._db.table(_TABLE).insert(payload).execute()
            rows = resp.data or []
            if rows:
                return rows[0]
        except Exception as exc:
            if strict:
                raise OptionalEvidencePersistError(str(exc)) from exc
            return payload
        if strict:
            raise OptionalEvidencePersistError("optional_evidence_submissions insert returned no rows")
        return payload

    def submit_file(
        self,
        *,
        user_id: str,
        proof_session_id: str | None,
        file_bytes: bytes,
        filename: str,
        source_type: SourceType = "document",
        extra_metadata: dict[str, Any] | None = None,
        strict: bool = False,
    ) -> dict[str, Any]:
        doc = extract_document_text(file_bytes, filename)
        analysis = analyze_from_extracted_document(doc, source_type=source_type)
        payload: dict[str, Any] = {
            "user_id": user_id,
            "proof_session_id": proof_session_id,
            "source_type": source_type,
            "status": analysis.status,
            "file_path": filename,
            "profile_url": None,
            "raw_text": None,
            "analysis_json": {
                **analysis.analysis_json,
                "file_name": doc.file_name,
                "file_type": doc.file_type,
                "extracted_text_preview": doc.preview,
                **(extra_metadata or {}),
            },
            "evidence_objects": analysis.evidence_objects,
        }
        return self._persist(payload, strict=strict)

    def submit_text(
        self,
        *,
        user_id: str,
        proof_session_id: str | None,
        source_type: SourceType,
        raw_text: str = "",
        profile_url: str | None = None,
        section_label: str | None = None,
        file_path: str | None = None,
        extra_metadata: dict[str, Any] | None = None,
        strict: bool = False,
    ) -> dict[str, Any]:
        analysis = analyze_optional_evidence(
            source_type=source_type,
            raw_text=raw_text,
            profile_url=profile_url,
            section_label=section_label,
            file_path=file_path,
        )
        payload = {
            "user_id": user_id,
            "proof_session_id": proof_session_id,
            "source_type": source_type,
            "status": analysis.status,
            "profile_url": profile_url,
            "file_path": file_path,
            "raw_text": raw_text or None,
            "analysis_json": {**analysis.analysis_json, **(extra_metadata or {})},
            "evidence_objects": analysis.evidence_objects,
        }
        return self._persist(payload, strict=strict)

    def list_for_session(self, *, user_id: str, proof_session_id: str) -> list[dict[str, Any]]:
        try:
            resp = (
                self._db.table(_TABLE)
                .select("*")
                .eq("user_id", user_id)
                .eq("proof_session_id", proof_session_id)
                .order("created_at", desc=True)
                .execute()
            )
            return resp.data or []
        except Exception:
            return []

    def list_standalone_for_user(
        self, *, user_id: str, source_types: tuple[str, ...] | None = None
    ) -> list[dict[str, Any]]:
        """List standalone (proof_session_id IS NULL) submissions for a user."""
        if isinstance(self._db, dict):
            rows = [
                row
                for row in self._db.get(_TABLE, {}).values()
                if str(row.get("user_id")) == user_id and row.get("proof_session_id") is None
            ]
            if source_types:
                rows = [r for r in rows if r.get("source_type") in source_types]
            rows.sort(key=lambda r: str(r.get("created_at") or ""), reverse=True)
            return rows
        try:
            query = (
                self._db.table(_TABLE)
                .select("*")
                .eq("user_id", user_id)
                .is_("proof_session_id", "null")
            )
            if source_types:
                query = query.in_("source_type", list(source_types))
            resp = query.order("created_at", desc=True).execute()
            return resp.data or []
        except Exception:
            return []

    def reextract_blocks_from_retained_original(
        self, *, user_id: str, evidence_id: str
    ) -> dict[str, Any] | None:
        """Deterministic block re-extraction from the RETAINED original file.

        Additive and idempotent: the original analysis and its evidence objects
        are preserved; newly extracted block-typed evidence (tables, charts,
        diagrams, code blocks, metrics — with section/block locators) is merged
        in, deduplicated by (skill, snippet). Nothing is inferred beyond the
        structural extraction; when no retained original exists the submission
        is returned unchanged with no fabricated blocks.
        """
        row = self.get_by_id(user_id=user_id, evidence_id=evidence_id)
        if not isinstance(row, dict):
            return None
        try:
            from app.services import proof_artifact_service

            artifacts = [
                a
                for a in proof_artifact_service.list_artifacts_for_proof(
                    self._db, proof_type="document", proof_id=evidence_id
                )
                if str(a.get("owner_user_id") or "") == str(user_id)
                and a.get("artifact_type") == "document_original"
                and a.get("retained")
            ]
        except Exception:
            artifacts = []
        if not artifacts:
            return row
        data = proof_artifact_service.fetch_artifact_bytes(self._db, artifacts[-1])
        if not data:
            return row
        filename = str(
            artifacts[-1].get("file_name")
            or row.get("file_path")
            or "document"
        )
        doc = extract_document_text(data, filename)
        if doc.status != "ok":
            return row
        analysis = analyze_from_extracted_document(doc, source_type="document")

        existing = [o for o in (row.get("evidence_objects") or []) if isinstance(o, dict)]
        seen = {
            (str(o.get("skill_name") or ""), str(o.get("snippet") or "")[:120])
            for o in existing
        }
        merged = list(existing)
        added = 0
        for obj in analysis.evidence_objects:
            key = (str(obj.get("skill_name") or ""), str(obj.get("snippet") or "")[:120])
            if key in seen:
                continue
            seen.add(key)
            merged.append(obj)
            added += 1
        # Upgrade pre-block duplicates in place: an existing keyword-only object
        # whose exact snippet now carries a block locator gains that locator
        # (same stored snippet, richer provenance — never overwritten otherwise).
        block_by_key = {
            (str(o.get("skill_name") or ""), str(o.get("snippet") or "")[:120]): o
            for o in analysis.evidence_objects
            if o.get("block_type")
        }
        for obj in existing:
            if obj.get("block_type"):
                continue
            match = block_by_key.get(
                (str(obj.get("skill_name") or ""), str(obj.get("snippet") or "")[:120])
            )
            if match:
                for field_name in (
                    "block_type", "block_index", "table_cells",
                    "visual_description", "nearby_caption", "figure_reference",
                    "section_label", "reason",
                ):
                    if match.get(field_name) is not None:
                        obj[field_name] = match[field_name]

        analysis_json = dict(row.get("analysis_json") or {})
        analysis_json["block_extraction_version"] = 2
        analysis_json["detected_skills"] = sorted(
            {str(o.get("skill_name") or "") for o in merged if o.get("skill_name")}
        )
        patch = {
            "evidence_objects": merged,
            "analysis_json": analysis_json,
            "updated_at": datetime.now(UTC).isoformat(),
        }
        if isinstance(self._db, dict):
            row.update(patch)
            return row
        try:
            resp = (
                self._db.table(_TABLE)
                .update(patch)
                .eq("id", evidence_id)
                .eq("user_id", user_id)
                .execute()
            )
            rows = resp.data or []
            return rows[0] if rows else {**row, **patch}
        except Exception as exc:
            raise OptionalEvidencePersistError(str(exc)) from exc

    def get_by_id(self, *, user_id: str, evidence_id: str) -> dict[str, Any] | None:
        if isinstance(self._db, dict):
            row = self._db.get(_TABLE, {}).get(evidence_id)
            if not row or str(row.get("user_id")) != user_id:
                return None
            return row
        try:
            resp = (
                self._db.table(_TABLE)
                .select("*")
                .eq("id", evidence_id)
                .eq("user_id", user_id)
                .limit(1)
                .execute()
            )
            rows = resp.data or []
            return rows[0] if rows else None
        except Exception:
            return None
